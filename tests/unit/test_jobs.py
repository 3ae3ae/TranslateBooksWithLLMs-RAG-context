import json
import sys
import threading
import types
from pathlib import Path
from unittest.mock import Mock

import pytest
from flask import Flask

from src.api.blueprints.jobs_routes import create_jobs_blueprint, job_summary
from src.api.job_runner import JobRunner
from src.api.translation_state import TranslationStateManager
from src.persistence.checkpoint_manager import CheckpointManager


@pytest.fixture
def jobs(tmp_path, monkeypatch):
    class ConnectionNotFound(Exception):
        pass
    connection = dict(connection_id='local', connection_type='openai_compatible', connection_name='Test',
                      llm_provider='openai', llm_api_endpoint='http://localhost:8000/v1/chat/completions',
                      api_key='YOUR_API_KEY_HERE', model='test', reasoning_effort='off', local=False,
                      request_timeout=17, max_attempts=3, temperature=0.2,
                      rate_limit_behavior='pause', rate_limit_retry_seconds=8)
    resolve = Mock(side_effect=lambda cid, model, reasoning: {**connection, 'connection_id': cid, 'model': model, 'reasoning_effort': reasoning})
    preferences = dict(default_connection_id='local', default_model='test', source_language='English', target_language='Korean', notes_mode='build', passage_tokens=123)
    remember = Mock()
    notes_calls = []
    def prepare(mode, **kw):
        notes_calls.append((mode, kw))
        return None if mode == 'off' else f'notes-{len(notes_calls)}.txt'
    monkeypatch.setitem(sys.modules, 'src.core.llm.connections', types.SimpleNamespace(resolve_for_pipeline=resolve, ConnectionNotFound=ConnectionNotFound))
    monkeypatch.setitem(sys.modules, 'src.utils.preferences', types.SimpleNamespace(get_preferences=lambda: preferences.copy(), remember_job_defaults=remember))
    monkeypatch.setitem(sys.modules, 'src.utils.notes_store', types.SimpleNamespace(prepare_job_notes=prepare))
    manager = CheckpointManager(str(tmp_path / 'jobs.db'))
    manager.uploads_dir = tmp_path / 'preserved'
    manager.uploads_dir.mkdir()
    state = TranslationStateManager(manager)
    output = tmp_path / 'output'
    uploads = output / 'uploads'
    uploads.mkdir(parents=True)
    files = []
    for i in range(3):
        path = uploads / f'book{i}.txt'
        path.write_text('A passage in English.')
        files.append(dict(file_path=str(path), input_filename=path.name))
    executed = []
    def execute(tid, config):
        executed.append(tid)
        state.set_translation_field(tid, 'status', 'completed')
    socket = Mock()
    runner = JobRunner(state, output, socket, execute=execute)
    app = Flask(__name__)
    app.register_blueprint(create_jobs_blueprint(state, output, socket))
    return types.SimpleNamespace(client=app.test_client(), app=app, runner=runner, state=state, db=manager.db,
                                 files=files, prefs=preferences, resolve=resolve, connection=connection,
                                 notes_calls=notes_calls, remember=remember, executed=executed, socket=socket,
                                 ConnectionNotFound=ConnectionNotFound, output=output)


def create(jobs, **kwargs):
    response = jobs.client.post('/api/jobs', json={'files': [jobs.files[0]], **kwargs})
    assert response.status_code == 201, response.json
    return response.json['jobs'][0]['id']


@pytest.mark.parametrize('file_type', ['txt', 'srt'])
@pytest.mark.parametrize('mode', [None, 'build', 'copy'])
def test_translation_always_updates_notes_every_passage(jobs, file_type, mode):
    path = Path(jobs.files[0]['file_path']).with_suffix('.' + file_type)
    path.write_text('1\n00:00:01,000 --> 00:00:02,000\nHello Alice.\n' if file_type == 'srt' else 'Hello Alice.')
    notes = {'keep_updating': False, 'update_interval': 'ignored'}
    if mode:
        notes['mode'] = mode
    if mode == 'copy':
        notes['source'] = 'existing.txt'
    tid = create(jobs, files=[{'file_path': str(path)}], notes=notes)
    config = jobs.db.get_job(tid)['config']
    assert config['notes']['mode'] == (mode or 'build')
    assert config['notes']['keep_updating'] is True
    assert config['notes']['update_interval'] == 1
    assert config['prompt_options']['auto_update_context'] is True
    assert config['prompt_options']['novel_context_update_interval'] == 1
    assert config['notes_file']


@pytest.mark.parametrize('task,mode,message', [
    ('translate', 'off', 'Translation requires notes.mode to be build or copy'),
    ('polish', 'build', 'Polish supports existing read-only notes or notes off'),
])
def test_notes_modes_reject_unsupported_task_combinations(jobs, task, mode, message):
    response = jobs.client.post('/api/jobs', json={'files': [jobs.files[0]], 'task': task, 'notes': {'mode': mode}})
    assert response.status_code == 400
    assert response.json['error'] == message
    assert not jobs.notes_calls
    assert not jobs.db.list_jobs()


@pytest.mark.parametrize('notes,mode', [({}, 'off'), ({'mode': 'off'}, 'off'), ({'source': 'existing.txt'}, 'copy'), ({'mode': 'copy', 'source': 'existing.txt'}, 'copy')])
def test_polish_keeps_notes_read_only_or_off(jobs, notes, mode):
    tid = create(jobs, task='polish', notes={**notes, 'keep_updating': True})
    config = jobs.db.get_job(tid)['config']
    assert config['notes']['mode'] == mode
    assert config['notes']['keep_updating'] is False
    assert config['prompt_options']['auto_update_context'] is False


def test_creation_freezes_defaults_without_keys(jobs):
    tid = create(jobs)
    config = jobs.db.get_job(tid)['config']
    assert config['max_tokens_per_chunk'] == 123
    assert config['request_timeout'] == 17
    assert config['parallel_workers'] == 1
    assert 'YOUR_API_KEY_HERE' not in json.dumps(config)
    jobs.prefs['passage_tokens'] = 999
    jobs.connection['request_timeout'] = 99
    observed = []
    def execute(tid, config):
        observed.append(config)
        jobs.state.set_translation_field(tid, 'status', 'completed')
    jobs.runner.execute = execute
    jobs.runner.run_next()
    assert observed[0]['connection']['request_timeout'] == 17
    assert observed[0]['connection']['api_key'] == 'YOUR_API_KEY_HERE'
    assert observed[0]['max_tokens_per_chunk'] == 123
    assert 'api_key' not in jobs.client.get(f'/api/jobs/{tid}').json['config']['connection']


def test_fifo_pause_move_and_resume_next(jobs):
    ids = [item['id'] for item in jobs.client.post('/api/jobs', json={'files': jobs.files}).json['jobs']]
    jobs.runner.pause_queue(True)
    assert not jobs.runner.run_next()
    jobs.runner.move(ids[2], 'top')
    jobs.runner.pause(ids[1])
    assert jobs.db.queue_state()['order'] == [ids[2], ids[0]]
    jobs.client.post(f'/api/jobs/{ids[1]}/resume')
    assert jobs.db.queue_state()['order'] == [ids[1], ids[2], ids[0]]
    jobs.runner.pause_queue(False)
    while jobs.runner.run_next():
        pass
    assert jobs.executed == [ids[1], ids[2], ids[0]]


def test_resume_does_not_preempt(jobs):
    first = create(jobs)
    second = create(jobs)
    entered, release = threading.Event(), threading.Event()
    def execute(tid, config):
        entered.set()
        assert release.wait(3)
        jobs.state.set_translation_field(tid, 'status', 'completed')
    jobs.runner.execute = execute
    worker = threading.Thread(target=jobs.runner.run_next)
    worker.start()
    assert entered.wait(3)
    jobs.runner.pause(second)
    jobs.client.post(f'/api/jobs/{second}/resume')
    assert jobs.runner.active == first
    assert jobs.db.queue_state()['order'] == [second]
    release.set()
    worker.join(3)
    assert not worker.is_alive()


def test_restart_holds_queue_and_marks_running(jobs):
    tid = create(jobs)
    jobs.db.update_job_progress(tid, status='running')
    jobs.runner.restart_hold()
    assert jobs.db.queue_state()['paused'] is True
    assert jobs.db.queue_state()['reason'] == 'restart'
    job = jobs.db.get_job(tid)
    assert job['status'] == 'interrupted'
    assert job['config']['status_reason'] == 'server_restart'
    assert not jobs.runner.run_next()


def test_resume_lifts_restart_hold_but_not_user_pause(jobs):
    tid = create(jobs)
    jobs.db.update_job_progress(tid, status='running')
    jobs.runner.restart_hold()
    assert jobs.client.post(f'/api/jobs/{tid}/resume', json={}).status_code == 200
    assert jobs.db.queue_state()['paused'] is False
    jobs.runner.pause(tid)
    jobs.runner.pause_queue(True)
    assert jobs.client.post(f'/api/jobs/{tid}/resume', json={}).status_code == 200
    assert jobs.db.queue_state()['paused'] is True


@pytest.mark.parametrize('stored,expected', [('interrupted','paused'), ('error','failed'), ('rate_limited','paused'), ('partial','partial')])
def test_status_mapping(jobs, stored, expected):
    tid = create(jobs)
    jobs.db.update_job_progress(tid, status=stored)
    assert job_summary(jobs.runner, jobs.db.get_job(tid))['status'] == expected


def test_elapsed_and_eta_survive_resume(jobs):
    tid = create(jobs)
    config = jobs.db.get_job(tid)['config']
    config['stats'] = {'elapsed_seconds': 60}
    jobs.db.update_job_config(tid, config)
    jobs.db.update_job_progress(tid, status='running')
    # After a resume the pipeline's own timer and ETA restart from zero.
    if jobs.state.get_translation(tid) is None:
        jobs.state.create_translation(tid, config)
    jobs.state.update_stats(tid, {'elapsed_seconds': 1, 'eta_seconds': 1, 'completed_chunks': 10, 'total_chunks': 20})
    summary = job_summary(jobs.runner, jobs.db.get_job(tid))
    assert summary['elapsed_seconds'] == 60
    assert summary['eta_seconds'] == 60


def test_series_waits_and_copies_notes_lazily(jobs):
    ids = [item['id'] for item in jobs.client.post('/api/jobs', json={'files': jobs.files, 'notes': {'mode': 'build', 'series': True}}).json['jobs']]
    assert len(jobs.notes_calls) == 1
    jobs.runner.pause(ids[0])
    assert not jobs.runner.run_next()
    jobs.client.post(f'/api/jobs/{ids[0]}/resume')
    assert jobs.runner.run_next()
    assert len(jobs.notes_calls) == 1
    assert jobs.runner.run_next()
    assert jobs.notes_calls[1][0] == 'copy'
    assert jobs.notes_calls[1][1]['carry_from'] == 'notes-1.txt'
    assert jobs.db.get_job(ids[1])['config']['notes_file'] == 'notes-2.txt'


def test_missing_connection_fails_without_execution(jobs):
    tid = create(jobs)
    jobs.resolve.side_effect = jobs.ConnectionNotFound()
    jobs.runner.run_next()
    job = jobs.client.get(f'/api/jobs/{tid}').json
    assert job['status'] == 'failed'
    assert job['status_reason'] == 'connection_removed'
    assert not jobs.executed


def test_passages_paging_search_preview_and_persisted_log(jobs):
    tid = create(jobs)
    for index in range(3):
        jobs.db.save_chunk(tid, index, f'[id0]source {index}', 'translation' + 'x' * 5000)
    result = jobs.client.get(f'/api/jobs/{tid}/passages?offset=1&limit=1&q=source').json
    assert result['total'] == 3
    assert result['items'][0]['source'] == 'source 1'
    assert len(jobs.client.get(f'/api/jobs/{tid}').json['latest_preview']['translation']) == 4000
    jobs.state.append_log(tid, {'message':'persist me'})
    jobs.state._translations.clear()
    assert jobs.client.get(f'/api/jobs/{tid}/log?tail=1').json == {'items':[{'message':'persist me'}]}


@pytest.mark.parametrize('delete_output', [False, True])
def test_delete_with_and_without_output(jobs, delete_output):
    tid = create(jobs)
    config = jobs.db.get_job(tid)['config']
    output = jobs.output / config['output_filename']
    output.write_text('output')
    jobs.db.save_chunk(tid, 0, 'source', 'translation')
    response = jobs.client.delete(f'/api/jobs/{tid}?delete_output={int(delete_output)}')
    assert response.status_code == 200
    assert output.exists() is not delete_output
    assert jobs.db.get_chunks(tid) == []
    assert tid not in jobs.db.queue_state()['order']


def test_continue_preserves_old_output_on_success(jobs):
    tid = create(jobs)
    jobs.runner.run_next()
    old = jobs.output / jobs.db.get_job(tid)['config']['output_filename']
    old.write_text('old output')
    response = jobs.client.post(f'/api/jobs/{tid}/continue', json=jobs.files[1])
    assert response.status_code == 201
    new_id = response.json['id']
    def execute(tid, config):
        path = jobs.output / 'temporary.txt'
        path.write_text('new output')
        jobs.state.set_translation_field(tid, 'output_filepath', str(path))
        jobs.state.set_translation_field(tid, 'status', 'completed')
    jobs.runner.execute = execute
    jobs.runner.run_next()
    assert old.read_text() == 'new output'
    assert old.with_name(f'{old.stem} (before update){old.suffix}').read_text() == 'old output'
    assert jobs.db.get_job(new_id)['config']['output_filepath'] == str(old)


def test_validation_is_atomic_and_rejects_paths(jobs):
    result = jobs.client.post('/api/jobs', json={'files':[jobs.files[0], {'file_path':'../outside.txt'}]})
    assert result.status_code == 400
    assert not jobs.db.list_jobs()
    assert not jobs.notes_calls


def test_upload_limit_loopback_and_sweep(tmp_path):
    import os
    from src.utils.security import RateLimiter
    from src.api.blueprints.security_routes import sweep_abandoned_uploads
    limiter = RateLimiter()
    assert all(limiter.is_allowed('127.0.0.1') for _ in range(30))
    assert all(limiter.is_allowed('::1') for _ in range(30))
    assert sum(limiter.is_allowed('192.0.2.1') for _ in range(30)) == 10
    upload = tmp_path / 'old.txt'
    upload.write_text('abandoned')
    os.utime(upload, (0,0))
    db = Mock(list_jobs=lambda: [])
    assert sweep_abandoned_uploads(tmp_path, db, now=90000) == 1
    assert not upload.exists()


@pytest.mark.asyncio
@pytest.mark.parametrize('file_type', ['txt', 'srt', 'epub', 'docx'])
async def test_connection_knobs_reach_client_factory(jobs, monkeypatch, file_type):
    from src.core.adapters.translate_file import translate_file
    connection = jobs.connection.copy()
    if file_type == 'epub':
        factory = Mock(return_value=None)
        monkeypatch.setattr('src.core.llm_client.create_llm_client', factory)
    elif file_type == 'docx':
        from unittest.mock import AsyncMock
        factory = Mock(return_value=object())
        monkeypatch.setattr('src.core.llm_client.LLMClient', factory)
        monkeypatch.setattr('src.core.docx.translator.translate_docx_file', AsyncMock(return_value={'success':True}))
    else:
        factory = Mock(side_effect=RuntimeError('stop after client creation'))
        monkeypatch.setattr('src.core.llm_client.LLMClient', factory)
    path = jobs.output / f'input.{file_type}'
    path.write_text('1\n00:00:01,000 --> 00:00:02,000\nHello world!\n' if file_type == 'srt' else 'A source passage.')
    jobs.state.get_checkpoint_manager().start_job('factory', file_type, {'file_type':file_type}, str(path))
    await translate_file(str(path), str(jobs.output / f'out.{file_type}'), 'English', 'Korean', 'test', 'openai',
                         jobs.state.get_checkpoint_manager(), 'factory', connection=connection,
                         max_tokens_per_chunk=123, parallel_workers=3, bilingual_output=True)
    passed = factory.call_args.kwargs['connection']
    assert passed == connection
    assert passed['request_timeout'] == 17
    assert passed['max_attempts'] == 3
    assert passed['temperature'] == 0.2


@pytest.mark.parametrize('file_type,fallback,failed', [('docx',0,1), ('epub',1,0), ('txt',0,1)])
def test_handler_classifies_failed_or_untranslated_output_as_partial(jobs, monkeypatch, file_type, fallback, failed):
    async def translate(**kwargs):
        Path(kwargs['output_filepath']).write_text('best available')
        kwargs['stats_callback']({'total_chunks':2, 'completed_chunks':1, 'failed_chunks':failed, 'fallback_used':fallback})
        assert kwargs['connection']['request_timeout'] == 17
        assert kwargs['max_attempts'] == 3
        assert kwargs['retry_delay'] == 2
        return failed == 0
    monkeypatch.setattr('src.api.handlers.translate_file', translate)
    tid = create(jobs)
    config = jobs.db.get_job(tid)['config']
    config.update(file_type=file_type, output_filename=f'output.{file_type}')
    jobs.db.update_job_config(tid, config)
    jobs.runner.execute = None
    jobs.runner.run_next()
    assert jobs.db.get_job(tid)['status'] == 'partial'


def test_handler_forwards_docx_bilingual(jobs, monkeypatch):
    from unittest.mock import AsyncMock
    from docx import Document
    from src.core.adapters.translate_file import translate_file
    import asyncio
    document = Document()
    document.add_paragraph('A paragraph.')
    source = jobs.output / 'source.docx'
    document.save(source)
    factory = Mock(return_value=object())
    translate = AsyncMock(return_value={'success':True})
    monkeypatch.setattr('src.core.llm.create_llm_provider', factory)
    monkeypatch.setattr('src.core.docx.translator.translate_docx_file', translate)
    asyncio.run(translate_file(str(source), str(jobs.output / 'out.docx'), 'English', 'Korean', 'test', 'openai',
                              jobs.state.get_checkpoint_manager(), 'docx', bilingual_output=True))
    assert translate.call_args.kwargs['bilingual'] is True
    assert translate.call_args.kwargs['prompt_options']['bilingual'] is True


def test_docx_persists_per_passage_rows_and_bilingual_output(jobs):
    import asyncio
    from docx import Document
    from src.core.docx.translator import translate_docx_file
    from tests.characterization.fake_llm import FakeEchoProvider
    source = jobs.output / 'book.docx'
    output = jobs.output / 'translated.docx'
    doc = Document()
    doc.add_paragraph('The first paragraph is here.')
    doc.add_paragraph('The second paragraph ends here.')
    doc.save(source)
    jobs.state.get_checkpoint_manager().start_job('docx-rows', 'docx', {'file_type':'docx'}, str(source))
    result = asyncio.run(translate_docx_file(str(source), str(output), 'English', 'Korean', 'test', FakeEchoProvider(),
                                           checkpoint_manager=jobs.state.get_checkpoint_manager(), translation_id='docx-rows', bilingual=True))
    assert result['success'] is True
    rows = jobs.db.get_chunks('docx-rows')
    assert rows and all(row['translated_text'] for row in rows)
    paragraphs = [p.text for p in Document(output).paragraphs]
    assert sum('The first paragraph is here.' in paragraph for paragraph in paragraphs) == 2


def test_notes_sections_dry_run_is_read_only(jobs, monkeypatch):
    from src.api.blueprints.translation_routes import create_translation_blueprint
    from src.utils.novel_context import build_novel_context, compress_dynamic_state
    tid = create(jobs)
    text = build_novel_context('# GLOBAL LORE', '## RELATIONSHIP EVOLUTION\n- A: friends.')
    for index in range(3):
        jobs.db.save_chunk(tid, index, 'source', 'translation', {'context_snapshot':compress_dynamic_state(text)})
    sections = [{'key':'relationships', 'title':'Relationships', 'scope':'story', 'editable':True, 'body':'friends'}]
    notes = sys.modules['src.utils.notes_store']  # the stub installed by the jobs fixture
    monkeypatch.setattr(notes, 'parse_sections', lambda _: sections, raising=False)
    monkeypatch.setattr(notes, 'serialize_sections', lambda _: text, raising=False)
    monkeypatch.setattr(notes, 'normalize_sections', lambda value: (value, ['duplicate']), raising=False)
    app = Flask(__name__)
    app.register_blueprint(create_translation_blueprint(jobs.state, lambda *_: None, jobs.output, jobs.socket))
    before = jobs.db.get_job(tid)
    chunks = jobs.db.get_chunks(tid)
    client = app.test_client()
    response = client.post(f'/api/translation/{tid}/context/0/resync', json={'sections':sections, 'dry_run':True, 'polish':True})
    assert response.status_code == 200, response.json
    assert response.json['affected_passages'] == 2
    assert response.json['polish_requests'] == 3
    assert response.json['normalized_sections'] == sections
    assert response.json['removed_lines'] == ['duplicate']
    assert jobs.db.get_job(tid) == before
    assert jobs.db.get_chunks(tid) == chunks
    get = client.get(f'/api/translation/{tid}/context/latest?format=sections')
    assert get.json['index'] == 2
    assert get.json['snapshot_indices'] == [0,1,2]


def test_gender_gating_stays_per_session(tmp_path, monkeypatch):
    import asyncio
    import src.config as config
    from src.utils.novel_context import open_novel_context_session
    from unittest.mock import AsyncMock
    monkeypatch.setattr(config, 'BYPASS_CONTEXT_GATING', True)
    (tmp_path / 'a.txt').write_text('# GLOBAL LORE')
    first = open_novel_context_session({'novel_context_file':'a.txt', 'bypass_context_gating':False}, tmp_path)
    second = open_novel_context_session({'novel_context_file':'a.txt', 'bypass_context_gating':True}, tmp_path)
    update = AsyncMock(return_value=('# GLOBAL LORE', '', []))
    monkeypatch.setattr('src.utils.novel_context.update_novel_context_chunk', update)
    asyncio.run(first.analyze_source(object(), 'test', 'source', 'English', 'Korean', 1, 1))
    assert update.call_args.kwargs['bypass_context_gating'] is False
    asyncio.run(second.analyze_source(object(), 'test', 'source', 'English', 'Korean', 1, 1))
    assert update.call_args.kwargs['bypass_context_gating'] is True
    assert config.BYPASS_CONTEXT_GATING is True


@pytest.mark.asyncio
@pytest.mark.parametrize('file_type', ['txt', 'srt', 'epub', 'docx'])
async def test_polish_parallel_bilingual_checkpoint_resume(jobs, monkeypatch, file_type):
    import asyncio
    from src.core.adapters.refine_file import refine_file
    from tests.characterization import fixtures
    from tests.characterization.fake_llm import FakeEchoProvider
    active, peak, calls = 0, 0, 0
    class Client(FakeEchoProvider):
        async def make_request(self, prompt, model=None, system_prompt=None):
            nonlocal active, peak, calls
            active += 1
            calls += 1
            peak = max(peak, active)
            try:
                await asyncio.sleep(0.001)
                return await self.generate(prompt, system_prompt=system_prompt)
            finally:
                active -= 1
        async def close(self):
            pass
    client = Client()
    factory = Mock(return_value=client)
    monkeypatch.setattr('src.core.translator.create_llm_client', factory)
    monkeypatch.setattr('src.core.llm_client.create_llm_client', factory)
    monkeypatch.setattr('src.core.refine.srt_refiner.create_llm_client', factory)
    source = getattr(fixtures, 'build_' + file_type)(jobs.output)
    output = jobs.output / ('polished.' + file_type)
    tid = 'polish-' + file_type
    jobs.state.get_checkpoint_manager().start_job(tid, file_type, {'file_type':file_type}, str(source))
    kwargs = dict(input_filepath=str(source), output_filepath=str(output), target_language='English', model_name='test',
                  llm_provider='openai', checkpoint_manager=jobs.state.get_checkpoint_manager(), translation_id=tid,
                  connection=jobs.connection.copy(), parallel_workers=3, bilingual_output=True, max_tokens_per_chunk=60,
                  srt_lines_per_block=2, auto_adjust_context=False)
    assert await refine_file(**kwargs)
    assert 1 < peak <= 3
    assert output.is_file()
    rows = jobs.db.get_chunks(tid)
    assert len(rows) > 1
    assert all(row['status'] == ('completed' if row['chunk_data']['polish_status'] == 'done' else 'failed') for row in rows)
    assert all(row['chunk_data']['draft_text'] == row['original_text'] for row in rows)
    failed_rows = [row for row in rows if row['status'] == 'failed']
    previous_calls = calls
    assert await refine_file(**kwargs)
    assert calls == previous_calls + len(failed_rows)
    assert factory.call_args.kwargs['connection']['request_timeout'] == 17



def test_queue_order_survives_another_runner(jobs):
    ids = [item['id'] for item in jobs.client.post('/api/jobs', json={'files':jobs.files}).json['jobs']]
    jobs.runner.move(ids[2], 'up')
    another = JobRunner(jobs.state, jobs.output, execute=jobs.runner.execute)
    assert another.db.queue_state()['order'] == [ids[0], ids[2], ids[1]]


def test_restart_pauses_a_queued_notes_rebuild(jobs):
    tid = create(jobs)
    jobs.runner.enqueue(tid, task=lambda: None)
    jobs.runner.restart_hold()
    job = jobs.db.get_job(tid)
    assert job['status'] == 'interrupted'
    assert job['config']['_context_resync']['status'] == 'paused'
    assert tid not in jobs.db.queue_state()['order']


def test_translate_only_notes_rebuild_can_queue_corrective_polish(jobs, monkeypatch):
    from src.api.blueprints.translation_routes import create_translation_blueprint
    from src.api.handlers import start_translation_job
    from src.utils.novel_context import build_novel_context, compress_dynamic_state
    tid = create(jobs)
    jobs.runner.run_next()
    config = jobs.db.get_job(tid)['config']
    original_source = config['preserved_input_path']
    output = jobs.output / config['output_filename']
    output.write_text('first translation')
    config['output_filepath'] = str(output)
    jobs.db.update_job_config(tid, config)
    jobs.state.set_translation_field(tid, 'output_filepath', str(output))
    text = build_novel_context('# GLOBAL LORE', '## RELATIONSHIP EVOLUTION\n- A: friends.')
    jobs.db.save_chunk(tid, 0, 'source', 'first translation', {'context_snapshot':compress_dynamic_state(text)})
    monkeypatch.setattr(jobs.runner, 'start', lambda: jobs.runner)
    def rebuild(tid, index, snapshot, socket, was_active, auto_resume, after, *args):
        config = jobs.db.get_job(tid)['config']
        config['_context_resync'].update(status='completed', last_processed_chunk=0)
        jobs.db.update_job_config(tid, config)
        assert after
        after()
    monkeypatch.setattr('src.core.adapters.generic_translator.resync_context_snapshots_background', rebuild)
    app = Flask(__name__)
    app.register_blueprint(create_translation_blueprint(jobs.state, lambda tid, cfg: start_translation_job(tid, cfg, jobs.state, jobs.output, jobs.socket), jobs.output, jobs.socket))
    response = app.test_client().post(f'/api/translation/{tid}/context/0/resync', json={'context_content':text, 'polish':True})
    assert response.status_code == 200, response.json
    preserved = jobs.db.get_job(tid)['config']['refinement_source_path']
    assert Path(preserved).read_text() == 'first translation'
    assert jobs.runner.run_next()
    assert jobs.db.queue_state()['order'] == [tid]
    def polish(tid, cfg):
        assert cfg['refine_only'] is True
        assert cfg['file_path'] == preserved
        output.write_text('corrected translation')
        jobs.state.set_translation_field(tid, 'status', 'completed')
    jobs.runner.execute = polish
    assert jobs.runner.run_next()
    config = jobs.db.get_job(tid)['config']
    assert config['refine_only'] is False
    assert config['preserved_input_path'] == original_source
    assert config['task'] == 'translate'
    assert Path(preserved).read_text() == 'first translation'
    assert output.read_text() == 'corrected translation'


def test_filename_pattern_uses_frozen_date_and_sanitized_model(jobs):
    tid = create(jobs, model='organization/model', output_filename_pattern='{date} {datetime} {model} {originalName}.{ext}')
    name = jobs.db.get_job(tid)['config']['output_filename']
    assert 'organization_model book0.txt' in name
    assert '{date}' not in name and '{datetime}' not in name


def test_log_tail_recovers_after_an_incomplete_write(jobs):
    tid = create(jobs)
    jobs.state.append_log(tid, {'message':'saved'})
    with jobs.state.log_path(tid).open('a') as stream:
        stream.write('{"message":')
    assert jobs.client.get(f'/api/jobs/{tid}/log').json['items'][-1] == {'message':'saved'}


def test_foreign_keys_and_orphan_cleanup(tmp_path):
    import sqlite3
    from src.persistence.database import Database
    path = tmp_path / 'old.db'
    with sqlite3.connect(path) as connection:
        connection.execute('CREATE TABLE checkpoint_chunks (translation_id TEXT, chunk_index INTEGER, original_text TEXT, translated_text TEXT, chunk_data JSON, status TEXT, completed_at TIMESTAMP, PRIMARY KEY (translation_id, chunk_index))')
        connection.execute("INSERT INTO checkpoint_chunks VALUES ('orphan', 0, 'source', NULL, NULL, 'failed', NULL)")
    db = Database(str(path))
    assert db._get_connection().execute('PRAGMA foreign_keys').fetchone()[0] == 1
    assert db.get_chunks('orphan') == []


def test_rate_limit_waiting_keeps_the_worker_slot(jobs, monkeypatch):
    from src.core.llm.exceptions import RateLimitError
    calls = []
    async def translate(**kwargs):
        calls.append(kwargs['translation_id'])
        if len(calls) == 1:
            assert jobs.runner.active == kwargs['translation_id']
            raise RateLimitError('test rate limit', retry_after=0.001)
        assert jobs.runner.active == kwargs['translation_id']
        Path(kwargs['output_filepath']).write_text('translation')
        kwargs['stats_callback']({'total_chunks':1, 'completed_chunks':1, 'failed_chunks':0})
        return True
    monkeypatch.setattr('src.api.handlers.translate_file', translate)
    tid = create(jobs)
    next_id = create(jobs)
    config = jobs.db.get_job(tid)['config']
    config.update(auto_pause_on_rate_limit=False, rate_limit_behavior='retry')
    jobs.db.update_job_config(tid, config)
    jobs.runner.execute = None
    assert jobs.runner.run_next()
    assert calls == [tid, tid]
    assert jobs.db.queue_state()['order'] == [next_id]
    assert jobs.client.get(f'/api/jobs/{tid}').json['status'] == 'completed'


def test_epub_partial_name_tracks_saved_output_and_clears_on_resume(jobs, monkeypatch):
    from src.utils.file_utils import get_partial_output_path
    calls = 0
    async def translate(**kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            path = get_partial_output_path(kwargs['output_filepath'])
            Path(path).write_text('partial output')
            jobs.state.set_interrupted(kwargs['translation_id'], True)
            return False
        assert not Path(kwargs['output_filepath']).name.startswith('[partial')
        Path(kwargs['output_filepath']).write_text('complete output')
        kwargs['stats_callback']({'total_chunks':1, 'completed_chunks':1, 'failed_chunks':0})
        return True
    monkeypatch.setattr('src.api.handlers.translate_file', translate)
    tid = create(jobs)
    config = jobs.db.get_job(tid)['config']
    config.update(file_type='epub', output_filename='book.epub')
    jobs.db.update_job_config(tid, config)
    jobs.runner.execute = None
    jobs.runner.run_next()
    summary = jobs.client.get(f'/api/jobs/{tid}').json
    assert summary['output_filename'] == '[partial] book.epub'
    assert summary['output_exists'] is True
    jobs.client.post(f'/api/jobs/{tid}/resume')
    jobs.runner.run_next()
    summary = jobs.client.get(f'/api/jobs/{tid}').json
    assert summary['status'] == 'completed'
    assert summary['output_filename'] == 'book.epub'
    assert summary['output_exists'] is True


def test_default_save_failure_does_not_hide_created_jobs(jobs):
    jobs.remember.side_effect = ValueError('defaults unavailable')
    tid = create(jobs)
    assert jobs.db.get_job(tid)
    assert jobs.client.get(f'/api/jobs/{tid}/log').json['items'][-1]['type'] == 'warning'


@pytest.mark.parametrize('change_model', [False, True])
def test_old_parallel_job_resumes_with_one_worker(jobs, change_model):
    tid = create(jobs)
    jobs.runner.pause(tid)
    config = jobs.db.get_job(tid)['config']
    config['parallel_workers'] = 8
    config['glossary_snapshot'] = {'Alice': 'old glossary'}
    config['prompt_options'].update(glossary_id=999, glossary_terms={'Alice': 'old glossary'})
    jobs.db.update_job_config(tid, config)
    body = {'model': 'replacement'} if change_model else {}
    response = jobs.client.post(f'/api/jobs/{tid}/resume', json=body)
    assert response.status_code == 200
    def execute(tid, config):
        assert config['parallel_workers'] == 1
        jobs.state.set_translation_field(tid, 'status', 'completed')
    jobs.runner.execute = execute
    assert jobs.runner.run_next()
    assert jobs.db.get_job(tid)['config']['parallel_workers'] == 1


def test_legacy_config_freezes_one_worker(jobs):
    from src.api.handlers import freeze_legacy_config
    config = {'model': 'test', 'connection': jobs.connection, 'parallel_workers': 8}
    frozen = freeze_legacy_config(config, jobs.state)
    assert frozen['parallel_workers'] == 1
    assert config['parallel_workers'] == 8


def test_passages_stage_contract_revision_and_socket_draft_final(jobs):
    tid = create(jobs, polish_after=True)
    source, draft, polished = '[id0]Source[id1]', '[id0]' + 'Draft ' * 900 + '[id1]', '[id0]Polished[id1]'
    revision = jobs.client.get(f'/api/jobs/{tid}').json['passages_revision']
    for polish_status, stage, text, status in [
        ('pending', 'draft', draft, 'partial'),
        ('failed', 'polish_failed', draft, 'failed'),
        ('done', 'polished', polished, 'completed'),
    ]:
        jobs.db.save_chunk(tid, 0, source, text,
                           {'draft_text': draft, 'polish_status': polish_status, 'passage_workflow_version': 1, 'local_to_global': {'0': 3}}, status=status)
        item = jobs.client.get(f'/api/jobs/{tid}/passages').json['items'][0]
        assert item == {'index': 0, 'source': 'Source', 'translation': text.replace('[id0]', '').replace('[id1]', ''),
                        'draft': 'Draft ' * 900, 'stage': stage, 'status': status}
        summary = jobs.client.get(f'/api/jobs/{tid}').json
        assert summary['passages_revision'] == revision + 1
        assert summary['current_unit'] is None
        emitted = [call.args[1] for call in jobs.socket.emit.call_args_list if call.args[0] == 'job_update'][-1]
        assert emitted['job']['passages_revision'] == revision + 1
        revision += 1
    jobs.db.save_chunk(tid, 0, source, polished, {'context_snapshot': 'snapshot'})
    metadata = jobs.db.get_chunks(tid)[0]['chunk_data']
    assert metadata['polish_status'] == 'done' and metadata['draft_text'] == draft and metadata['local_to_global'] == {'0': 3}
    listing = jobs.client.get('/api/jobs').json['jobs']
    assert next(item for item in listing if item['id'] == tid)['passages_revision'] == revision + 1


def test_translate_only_passages_have_null_draft_and_stage(jobs):
    tid = create(jobs)
    jobs.db.save_chunk(tid, 0, 'Source', 'Translation')
    assert jobs.client.get(f'/api/jobs/{tid}/passages').json['items'] == [
        {'index': 0, 'source': 'Source', 'translation': 'Translation', 'draft': None, 'stage': None, 'status': 'completed'}]


def test_legacy_refine_after_passages_are_exposed_as_drafts(jobs):
    tid = create(jobs, polish_after=True)
    jobs.db.save_chunk(tid, 0, 'Source', 'Legacy draft')
    item = jobs.client.get(f'/api/jobs/{tid}/passages').json['items'][0]
    assert item['stage'] == 'draft' and item['draft'] == 'Legacy draft'
