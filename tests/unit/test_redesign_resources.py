"""Resource contracts, isolated from credentials, provider APIs, and user files."""
import json
from pathlib import Path
import sqlite3
import stat
import sys
from types import ModuleType
from unittest.mock import Mock

from flask import Flask
import pytest
import requests

from src import config
from src.api.blueprints.estimate_routes import create_estimate_blueprint
from src.api.blueprints.notes_routes import create_notes_blueprint
from src.api.blueprints.preferences_routes import create_preferences_blueprint
from src.api.blueprints.styles_routes import create_styles_blueprint
from src.core.pricing import job_estimator
from src.utils import custom_instructions as styles, notes_store as notes, preferences


@pytest.fixture
def resources(monkeypatch):
    connection = dict(connection_id="test", llm_provider="openai", model="test-model", reasoning_effort="off", connection_type="openai", local=False, request_timeout=30, api_key="", llm_api_endpoint="https://example.test/v1/chat/completions")
    connections = ModuleType("src.core.llm.connections")
    connections.ConnectionNotFound = type("ConnectionNotFound", (ValueError,), {})
    default = {"connection_id": "test", "model": "test-model"}
    connections.get_default = Mock(side_effect=lambda: dict(default))

    def set_default(connection_id, model=None):
        default.update(connection_id=connection_id, model=model or default["model"])
        return dict(default)

    connections.set_default = Mock(side_effect=set_default)
    connections.resolve_for_pipeline = Mock(side_effect=lambda cid, model, reasoning="off": {**connection, "connection_id": cid, "model": model or connection["model"], "reasoning_effort": reasoning})
    connections.list_models = Mock(return_value=[{"id": "test-model", "pricing": {"input_per_m": 2, "output_per_m": 4}}])
    monkeypatch.setitem(sys.modules, "src.core.llm.connections", connections)
    app = Flask(__name__)
    app.register_blueprint(create_preferences_blueprint())
    app.register_blueprint(create_notes_blueprint())
    app.register_blueprint(create_styles_blueprint())
    app.register_blueprint(create_estimate_blueprint())
    return app.test_client(), connections


def seed_job(job_id, status, job_config):
    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(config.DATA_DIR / "jobs.db") as db:
        db.execute("CREATE TABLE IF NOT EXISTS translation_jobs (translation_id TEXT, status TEXT, config TEXT)")
        db.execute("INSERT INTO translation_jobs VALUES (?, ?, ?)", (job_id, status, json.dumps(job_config)))


def test_preferences_seed_schema_permissions_and_connection_authority(resources, monkeypatch):
    client, connections = resources
    monkeypatch.setattr(config, "MAX_TOKENS_PER_CHUNK", 700)
    result = client.get("/api/preferences").get_json()
    assert result["values"]["passage_tokens"] == 700
    assert result["values"]["default_connection_id"] == "test"
    assert result["values"]["default_model"] == "test-model"
    assert result["schema"]["passage_tokens"]["min"] == 50
    assert result["schema"]["default_reasoning"]["choices"][0] == "off"
    assert stat.S_IMODE(preferences.preferences_path().stat().st_mode) == 0o600
    monkeypatch.setattr(config, "MAX_TOKENS_PER_CHUNK", 900)
    assert preferences.get_preferences()["passage_tokens"] == 700
    connections.set_default.assert_not_called()


def test_preferences_validation_is_atomic(resources):
    client, _ = resources
    before = client.get("/api/preferences").get_json()["values"]
    result = client.patch("/api/preferences", json={"passage_tokens": 2, "bilingual": 1, "notes_update_interval": 0, "task": "other", "unknown": True})
    assert result.status_code == 400
    assert set(result.get_json()["error"]["details"]) == {"passage_tokens", "bilingual", "notes_update_interval", "task", "unknown"}
    assert client.get("/api/preferences").get_json()["values"] == before


@pytest.mark.parametrize("patch", [
    {"notes_mode": "off"},
    {"price_overrides": {"test:model": {"input_per_m": -1, "output_per_m": 1}}},
    {"price_overrides": {"test:model": {"input_per_m": True, "output_per_m": 1}}},
])
def test_preferences_nested_validation(resources, patch):
    assert resources[0].patch("/api/preferences", json=patch).status_code == 400


def test_remember_defaults_copies_job_fields(resources):
    _, connections = resources
    preferences.remember_job_defaults({"connection_id": "other", "model": "model", "reasoning": "low", "bilingual": True, "notes": {"mode": "copy", "update_interval": 3}, "files": []})
    values = preferences.get_preferences()
    assert (values["default_connection_id"], values["default_model"], values["default_reasoning"]) == ("other", "model", "low")
    assert values["notes_mode"] == "copy"
    assert "notes_update_interval" not in values
    assert values["bilingual"] is True
    connections.set_default.assert_called_with("other", "model")


def test_notes_round_trip_dry_run_matches_saved_normalization(resources):
    client, _ = resources
    assert client.post("/api/notes", json={"name": "book.txt"}).status_code == 201
    sections = client.get("/api/notes/book.txt").get_json()["sections"]
    assert [s["key"] for s in sections] == ["characters", "aliases", "name_map", "glossary", "addressing", "relationships"]
    assert notes.parse_sections(notes.serialize_sections(sections)) == sections
    sections[0]["body"] = "- 爱丽丝: female; explorer\n- 爱丽丝: female; explorer\n- None: unknown"
    sections[3]["body"] = "- 爱丽丝: Alice"
    path = config.NOVEL_CONTEXTS_DIR / "book.txt"
    before = path.read_bytes()
    dry = client.put("/api/notes/book.txt?dry_run=1", json={"sections": sections})
    assert dry.status_code == 200
    assert path.read_bytes() == before
    saved = client.put("/api/notes/book.txt", json={"sections": sections})
    assert saved.get_json() == dry.get_json()
    assert client.get("/api/notes/book.txt").get_json()["sections"] == dry.get_json()["normalized_sections"]
    assert dry.get_json()["removed_lines"]
    assert "- 爱丽丝: Alice" in dry.get_json()["normalized_sections"][2]["body"]


def test_notes_dry_run_does_not_report_reformatted_lines_as_removed():
    sections = notes.parse_sections(notes.serialize_sections([]))
    sections[0]["body"] = "- Carol: Female, librarian, scolds Alice for being late"
    normalized, removed = notes.normalize_sections(sections)
    assert "Carol" in normalized[0]["body"]
    assert removed == []


def test_notes_read_only_map_and_section_injection(resources):
    client, _ = resources
    notes.create_notes("book.txt")
    assert client.put("/api/notes/book.txt", json={"sections": [{"key": "name_map", "body": "- edited"}]}).status_code == 400
    assert client.put("/api/notes/book.txt", json={"sections": [{"key": "characters", "body": "## NAME TRANSLATION MAP\n- edited"}]}).status_code == 400
    assert client.put("/api/notes/book.txt", json={"text": "raw notes"}).status_code == 400


@pytest.mark.parametrize("status", ["running", "queued", "interrupted", "partial", "error"])
def test_notes_in_use_blocks_all_mutations(resources, status):
    client, _ = resources
    notes.create_notes("book.txt")
    seed_job("unfinished", status, {"notes_file": "book.txt"})
    assert notes.notes_in_use("book.txt") == ["unfinished"]
    assert client.get("/api/notes").get_json()["notes"][0]["used_by"] == ["unfinished"]
    for response in [client.put("/api/notes/book.txt?dry_run=1", json={"sections": []}), client.post("/api/notes/book.txt/rename", json={"name": "other.txt"}), client.delete("/api/notes/book.txt")]:
        assert response.status_code == 409
        assert response.get_json()["error"]["details"]["used_by"] == ["unfinished"]


def test_prepare_notes_build_copy_carry_unique_and_completed_mutation(resources):
    client, _ = resources
    first = notes.prepare_job_notes("build", job_title="book.epub")
    second = notes.prepare_job_notes("build", job_title="book.epub")
    assert first == "book_notes.txt" and second == "book_notes_2.txt"
    notes.put_notes(first, [{"key": "characters", "body": "- Alice: female; explorer"}])
    copied = notes.prepare_job_notes("copy", source=first, name="copied", job_title="book")
    carried = notes.prepare_job_notes("copy", source=second, carry_from=copied, name="copied", job_title="volume2")
    assert copied != carried
    assert notes.get_notes(carried)["sections"] == notes.get_notes(first)["sections"]
    assert notes.prepare_job_notes("off", job_title="book") is None
    assert notes.prepare_job_notes("build", name="My notes", job_title="book") == "My_notes.txt"
    seed_job("finished", "completed", {"notes_file": first})
    assert client.post(f"/api/notes/{first}/rename", json={"name": "renamed.txt"}).status_code == 200
    assert client.delete("/api/notes/renamed.txt").status_code == 200


def test_notes_and_styles_reject_traversal_and_external_symlinks(resources, tmp_path):
    client, _ = resources
    assert client.post("/api/notes", json={"name": "../outside.txt"}).status_code == 400
    assert client.post("/api/styles", json={"name": "../outside", "translation": "text"}).status_code == 400
    outside = tmp_path / "outside.txt"
    outside.write_text("outside")
    config.NOVEL_CONTEXTS_DIR.mkdir(parents=True)
    (config.NOVEL_CONTEXTS_DIR / "link.txt").symlink_to(outside)
    assert client.get("/api/notes/link.txt").status_code == 400
    config.CUSTOM_INSTRUCTIONS_DIR.mkdir(parents=True)
    (config.CUSTOM_INSTRUCTIONS_DIR / "link.txt").symlink_to(outside)
    assert client.get("/api/styles/link.txt").status_code == 400


def test_styles_crud_duplicate_and_legacy_save_conversion(resources):
    client, _ = resources
    result = client.post("/api/styles", json={"name": "noir", "translation": "Translate noir", "refinement": "Polish noir"})
    assert result.status_code == 201 and result.get_json()["name"] == "noir.yaml"
    assert client.get("/api/styles").get_json()["styles"][0]["name"] == "noir.yaml"
    assert client.post("/api/styles/noir.yaml/duplicate", json={"name": "copy"}).status_code == 201
    assert client.post("/api/styles", json={"name": "noir"}).status_code == 409
    assert client.put("/api/styles/noir", json={"translation": "Changed"}).get_json()["translation"] == "Changed"
    legacy = config.CUSTOM_INSTRUCTIONS_DIR / "legacy.txt"
    legacy.write_text("Legacy text")
    assert client.get("/api/styles/legacy.txt").get_json()["refinement"] == "Legacy text"
    result = client.put("/api/styles/legacy.txt", json={"translation": "New", "refinement": "Polished"})
    assert result.get_json()["name"] == "legacy.yaml" and not legacy.exists()
    assert client.delete("/api/styles/legacy.yaml").status_code == 200


def test_compare_uses_runtime_styles_and_connection_factory(resources, monkeypatch):
    from src.api.blueprints import sample_routes as sample
    styles.save_style("noir", {"translation": "Runtime translation", "refinement": "Runtime polish"}, create=True)
    options = sample._column_prompt_options({}, {"style": "noir"})
    assert options == {"custom_instructions": "Runtime translation", "refinement_instructions": "Runtime polish"}
    factory = Mock()
    monkeypatch.setattr(sample, "create_llm_provider", factory)
    sample._instantiate_provider({"connection_id": "test", "model": "test-model", "reasoning": "low"})
    assert factory.call_args.kwargs["connection"]["reasoning_effort"] == "low"
    assert sample._sample_max_chars({"passage_tokens": 300}) == 1200
    assert sample._sample_max_chars({"passage_tokens": 1000}) == 2000


def test_compare_snapshots_resources_and_changes_revision(resources, monkeypatch):
    from src.api.blueprints import sample_routes as sample
    from src.api.sample_state import SampleStateManager
    state = SampleStateManager()
    app = Flask("compare")
    app.register_blueprint(sample.create_sample_blueprint(state))
    client = app.test_client()
    uploads = config.INCOMING_DIR
    uploads.mkdir(parents=True)
    path = uploads / "book.txt"
    path.write_text("This is sample source.")
    monkeypatch.setattr(sample, "_load_source_units", lambda *a, **kw: [{"main_content": "First."}, {"main_content": "Middle."}, {"main_content": "Last."}])
    styles.save_style("noir", {"translation": "First style"}, create=True)
    notes.create_notes("book.txt")
    payload = dict(file_path=str(path), file_type="txt", source_language="English", target_language="Korean", mode="refine", defer_dispatch=True, notes="book.txt", glossary_id=1, columns=[{"connection_id": "test", "model": "test-model", "reasoning": "off", "style": "noir"}])
    response = client.post("/api/sample/run", json=payload)
    assert response.status_code == 200
    first = response.get_json()
    context = state.get_run_context(first["sample_id"])
    assert context["target_language"] == "Korean"
    assert context["columns"][0]["_prompt_options"]["custom_instructions"] == "First style"
    assert "_connection" not in first["columns"][0]
    styles.save_style("noir", {"translation": "Second style"})
    second = client.post("/api/sample/run", json=payload).get_json()
    assert first["columns"][0]["cache_key"] != second["columns"][0]["cache_key"]
    notes.put_notes("book.txt", [{"key": "characters", "body": "- Alice: female; explorer"}])
    third = client.post("/api/sample/run", json=payload).get_json()
    assert third["columns"][0]["cache_key"] != second["columns"][0]["cache_key"]


def test_estimate_requests_pricing_and_real_chunker(resources):
    client, connections = resources
    from pathlib import Path
    uploads = config.INCOMING_DIR
    uploads.mkdir(parents=True)
    path = uploads / "book.txt"
    path.write_text("\n\n".join(["A passage about Alice. " * 15] * 10))
    payload = dict(files=[{"file_path": str(path), "input_filename": "book.txt"}], connection_id="test", model="test-model", passage_tokens=100, polish_after=True, notes={"mode": "build", "keep_updating": False, "update_interval": "ignored", "consolidation_interval": 2, "prompt_max_tokens": 100, "source_memory_chars": 0})
    seed_job("history", "completed", {"connection_id": "test", "model": "test-model", "stats": {"requests": 10, "elapsed_seconds": 50}})
    result = client.post("/api/estimate", json=payload)
    assert result.status_code == 200
    value = result.get_json()
    from src.api.blueprints.sample_routes import _load_source_units
    passages = len(_load_source_units(str(path), "txt", passage_tokens=100))
    updates = passages
    assert value["passages"] == passages
    consolidations = sum(i % 2 == 0 or i == passages for i in range(1, passages + 1))
    assert value["requests"] == passages * 2 + updates + consolidations
    assert value["time_min_s"] == round(value["requests"] * 5 * 0.8)
    assert value["time_max_s"] == round(value["requests"] * 5 * 1.2)
    assert value["time_source"] == "history" and value["price_source"] == "models"
    assert value["cost_max_usd"] >= value["cost_min_usd"] > 0
    assert value["billing"] == "api"
    preferences.patch_preferences({"price_overrides": {"test:test-model": {"input_per_m": 1, "output_per_m": 2}}})
    assert job_estimator.resolve_pricing(connections.resolve_for_pipeline("test", "test-model"), preferences.get_preferences()["price_overrides"])[1] == "models"
    connections.list_models.return_value = []
    assert client.post("/api/estimate", json=payload).get_json()["price_source"] == "override"


@pytest.mark.parametrize("local,kind,billing", [(True, "ollama", "local"), (False, "ollama_cloud", "plan"), (False, "openai", "api"), (False, "neuralwatt", "api")])
def test_estimate_srt_lines_read_only_notes_and_billing(resources, local, kind, billing):
    client, connections = resources
    from pathlib import Path
    uploads = config.INCOMING_DIR
    uploads.mkdir(parents=True)
    path = uploads / "subtitles.srt"
    path.write_text("\n\n".join(f"{i}\n00:00:01,000 --> 00:00:02,000\nHello Alice." for i in range(1, 6)))
    connections.resolve_for_pipeline.side_effect = lambda *a: dict(connection_id="test", model="unknown", llm_provider="openai", local=local, connection_type=kind)
    connections.list_models.return_value = []
    payload = dict(files=[{"file_path": str(path)}], connection_id="test", model="unknown", task="polish", srt_lines_per_block=2, notes={"mode": "copy", "keep_updating": False})
    value = client.post("/api/estimate", json=payload).get_json()
    assert value["passages"] == value["requests"] == 3
    assert value["billing"] == billing and value["time_source"] == "unknown"
    assert value["cost_min_usd"] == (None if billing == "api" else 0)
    payload["files"][0]["file_path"] = "../../outside.txt"
    assert client.post("/api/estimate", json=payload).status_code == 400


@pytest.mark.parametrize("simplify", [False, True])
def test_estimate_rich_formats_use_pipeline_chunkers(resources, tmp_path, sampler_epub_path, simplify):
    from docx import Document
    from src.core.common.plain_text_pipeline import build_plain_segments
    from src.core.docx.docx_translation_adapter import DocxTranslationAdapter
    from src.core.docx.plain_extractor import extract_plain_paragraphs
    path = tmp_path / "book.docx"
    document = Document()
    document.add_heading("Chapter One", 1)
    for _ in range(10):
        document.add_paragraph("Alice explored the mountain. " * 15)
    document.save(path)
    units = job_estimator.load_job_units(path, "docx", passage_tokens=150, srt_lines_per_block=10, chapter_mode=True, simplify_formatting=simplify)
    if simplify:
        content = extract_plain_paragraphs(str(path))
        expected = build_plain_segments(content.paragraphs_text, 150, paragraph_kinds=content.paragraphs_style, chapter_mode=True)
    else:
        adapter = DocxTranslationAdapter()
        html, context = adapter.extract_content(str(path), None)
        text, tags, _ = adapter.preserve_structure(html, context, None)
        expected = adapter.create_chunks(text, tags, 150, None, chapter_mode=True)
    assert len(units) == len(expected) > 1
    epub_units = job_estimator.load_job_units(sampler_epub_path, "epub", passage_tokens=150, srt_lines_per_block=10, chapter_mode=True, simplify_formatting=simplify)
    assert len(epub_units) > 10
    assert all(unit["main_content"] for unit in epub_units)


def test_preferences_ignore_removed_saved_keys(resources):
    client, _ = resources
    preferences._atomic_write({"notifications": {"enabled": True}, "unknown": True, "bilingual": True})
    result = client.get("/api/preferences")
    assert result.status_code == 200
    assert result.json["values"]["bilingual"] is True
    assert "notifications" not in result.json["values"]
    assert "unknown" not in result.json["values"]


def test_notification_test_route_is_removed(resources):
    from src.api.blueprints.config_routes import create_config_blueprint
    app = Flask("removed-notifications")
    app.register_blueprint(create_config_blueprint())
    assert app.test_client().post("/api/notifications/test").status_code == 404


def test_folder_routes_use_runtime_paths(resources, monkeypatch):
    from src.api.blueprints.file_routes import create_file_blueprint
    from src.api.services import FileService
    opened = []

    def open_folder(path):
        opened.append(Path(path))
        return True, "Opened", str(path)

    monkeypatch.setattr(FileService, "open_system_folder", staticmethod(open_folder))
    app = Flask("folders")
    app.register_blueprint(create_file_blueprint(config.OUTPUT_DIR))
    client = app.test_client()
    for folder in ("output", "notes", "styles"):
        assert client.post(f"/api/folders/{folder}/open").status_code == 200
    assert opened == [Path(config.OUTPUT_DIR), config.NOVEL_CONTEXTS_DIR, config.CUSTOM_INSTRUCTIONS_DIR]
    assert client.post("/api/folders/unknown/open").status_code == 404


def test_old_notes_preferences_fall_back_to_build(resources):
    client, _ = resources
    preferences._atomic_write({"notes_mode": "off", "notes_update_interval": 9})
    result = client.get("/api/preferences").json
    assert result["values"]["notes_mode"] == "build"
    assert result["schema"]["notes_mode"]["choices"] == ["build", "copy"]
    assert "notes_update_interval" not in result["values"]
    assert "notes_update_interval" not in result["schema"]


def test_remember_polish_without_notes_keeps_translation_notes_default(resources):
    preferences.remember_job_defaults({"task": "polish", "notes": {"mode": "off"}})
    assert preferences.get_preferences()["notes_mode"] == "build"
    assert preferences.get_preferences()["task"] == "polish"


@pytest.mark.parametrize("task,mode,message", [
    ("translate", "off", "Translation requires notes.mode to be build or copy"),
    ("polish", "build", "Polish supports existing read-only notes or notes off"),
])
def test_estimate_rejects_unsupported_notes_modes(resources, task, mode, message):
    response = resources[0].post("/api/estimate", json={"files": [{"file_path": "unused"}], "task": task, "notes": {"mode": mode}})
    assert response.status_code == 400
    assert response.json["error"]["message"] == message


@pytest.mark.parametrize("mode", [None, "copy"])
def test_estimate_srt_translation_updates_notes_for_every_passage(resources, mode):
    client, _ = resources
    uploads = config.INCOMING_DIR
    uploads.mkdir(parents=True)
    path = uploads / "subtitles.srt"
    path.write_text("\n\n".join(f"{i}\n00:00:01,000 --> 00:00:02,000\nHello Alice." for i in range(1, 6)))
    note_options = {"keep_updating": False, "update_interval": None, "consolidation_interval": 2}
    if mode:
        notes.create_notes("existing.txt")
        note_options.update(mode=mode, source="existing.txt")
    response = client.post("/api/estimate", json={"files": [{"file_path": str(path)}], "srt_lines_per_block": 2, "notes": note_options})
    assert response.status_code == 200
    file = response.json["per_file"][0]
    assert file["passages"] == file["notes_updates"] == 3
    assert file["consolidation_calls"] == 2
    assert file["requests"] == 8


def test_interleaved_estimate_adds_one_polish_request_without_extra_notes(resources):
    client, _ = resources
    uploads = config.INCOMING_DIR
    uploads.mkdir(parents=True)
    path = uploads / 'chapter.txt'
    path.write_text('A passage about Alice. ' * 100)
    payload = dict(files=[{'file_path': str(path)}], connection_id='test', model='test-model',
                   passage_tokens=60, notes={'mode': 'build', 'consolidation_interval': 2})
    translated = client.post('/api/estimate', json={**payload, 'polish_after': False}).json
    polished = client.post('/api/estimate', json={**payload, 'polish_after': True}).json
    assert polished['requests'] - translated['requests'] == translated['passages']
    for key in ('notes_updates', 'consolidation_calls'):
        assert polished['per_file'][0][key] == translated['per_file'][0][key]
    assert polished['input_tokens'] > translated['input_tokens']
