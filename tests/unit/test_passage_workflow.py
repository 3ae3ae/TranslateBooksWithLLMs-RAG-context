"""Offline coverage of durable translation and polish at source boundaries."""
import re
import zipfile

import pytest
from docx import Document
from lxml import etree

from src.config import TRANSLATE_TAG_IN, TRANSLATE_TAG_OUT
from src.core.adapters.translate_file import translate_file
from src.core.common.passage_workflow import PassageWorkflow, is_polished
from src.core.llm.base import LLMResponse
from src.core.llm.exceptions import RateLimitError
from tests.characterization import fake_llm, fixtures
from tests.characterization.recorder import _checkpoint_manager


@pytest.fixture
def provider(monkeypatch):
    fake_llm.install(monkeypatch)
    calls = []
    control = {'fail_polish': False, 'fail_translation': False, 'rate_limit': False, 'corrupt': False}

    async def generate(self, prompt, timeout=0, system_prompt=None):
        text = fake_llm._extract_source_block(prompt)
        stage = 'P' if 'REFIN' in prompt.split('<SOURCE_TEXT>')[0].upper() else 'T'
        calls.append((stage, tuple(re.findall(r'(?:SOURCE|DRAFT)(\d+)', text)), prompt))
        if stage == 'P' and control['rate_limit']:
            raise RateLimitError('offline limit', provider='fake')
        if (stage == 'P' and control['fail_polish']) or (stage == 'T' and control['fail_translation']):
            return None
        text = text.replace('DRAFT', 'POLISHED') if stage == 'P' else text.replace('SOURCE', 'DRAFT')
        if stage == 'P' and control['corrupt']:
            text = re.sub(r'\[id\d+\]', '', text)
        return LLMResponse(content=f'{TRANSLATE_TAG_IN}{text}{TRANSLATE_TAG_OUT}',
                           prompt_tokens=10, completion_tokens=10, context_used=20, context_limit=4096)

    monkeypatch.setattr(fake_llm.FakeEchoProvider, 'generate', generate)
    return calls, control


def source_file(tmp_path, fmt):
    paragraphs = [f'SOURCE{i} ' + 'A traveller walked through the quiet village. ' * 3 for i in range(4)]
    path = tmp_path / f'book.{fmt}'
    if fmt == 'txt':
        path.write_text('\n\n'.join(paragraphs))
    elif fmt == 'srt':
        path.write_text('\n\n'.join(f'{i * 3}\n00:00:0{i},000 --> 00:00:0{i + 1},000\n{text}' for i, text in enumerate(paragraphs)))
    elif fmt == 'docx':
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(path)
    else:
        template = fixtures.build_epub(tmp_path)
        with zipfile.ZipFile(template) as old, zipfile.ZipFile(path, 'w') as new:
            for info in old.infolist():
                data = old.read(info)
                if info.filename.endswith(('.xhtml', '.html')):
                    data = ('<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Book</title></head><body>'
                            + ''.join(f'<h1>{text}</h1>' if i == 0 else f'<p><em>{text}</em></p>' for i, text in enumerate(paragraphs))
                            + '</body></html>').encode()
                new.writestr(info, data)
    return path


def setup_job(tmp_path, fmt, **options):
    manager = _checkpoint_manager(tmp_path)
    source = source_file(tmp_path, fmt)
    output = tmp_path / f'out.{fmt}'
    manager.start_job('job', fmt, {'refine_after': True, '_passage_workflow': 1}, str(source))
    prompt_options = {'_passage_workflow': 1, '_draft_artifacts': {}, **options}
    kwargs = dict(input_filepath=str(source), output_filepath=str(output), source_language='English',
                  target_language='Korean', model_name='fake-echo', llm_provider='poe', checkpoint_manager=manager,
                  translation_id='job', max_tokens_per_chunk=60, prompt_options=prompt_options,
                  srt_lines_per_block=2, parallel_workers=3, auto_adjust_context=False)
    return manager, output, kwargs


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt,options', [('txt', {}), ('srt', {}), ('epub', {}),
    ('epub', {'chapter_mode': True}), ('epub', {'plain_text_mode': True}),
    ('docx', {}), ('docx', {'plain_text_mode': True})])
async def test_interleaved_formats_keep_source_draft_final_and_order(tmp_path, provider, fmt, options):
    calls, _ = provider
    manager, output, kwargs = setup_job(tmp_path, fmt, **options)
    assert await translate_file(**kwargs)
    rows = manager.db.get_chunks('job')
    assert len(rows) > 1
    assert [call[0] for call in calls] == ['T', 'P'] * len(rows)
    assert [call[1] for call in calls[::2]] == [call[1] for call in calls[1::2]]
    assert all('SOURCE' in row['original_text'] and 'DRAFT' in row['chunk_data']['draft_text']
               and 'POLISHED' in row['translated_text'] and is_polished(row) for row in rows)
    artifact = kwargs['prompt_options']['_draft_artifacts']['output']
    if fmt in ('txt', 'srt'):
        assert b'DRAFT' in artifact and b'POLISHED' not in artifact
        assert 'POLISHED' in output.read_text()
    else:
        import io
        with zipfile.ZipFile(io.BytesIO(artifact)) as archive:
            assert any(b'DRAFT' in archive.read(name) for name in archive.namelist())
            assert not any(b'POLISHED' in archive.read(name) for name in archive.namelist())
    if fmt == 'srt':
        text = output.read_text()
        assert [line for line in text.splitlines() if '-->' in line] == [line for line in source_file(tmp_path, fmt).read_text().splitlines() if '-->' in line]
        assert [line for line in text.splitlines() if line.isdigit()] == ['0', '3', '6', '9']
    calls.clear()
    assert await translate_file(**kwargs)
    assert calls == []


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt,options', [('txt', {}), ('srt', {}), ('epub', {}),
    ('epub', {'plain_text_mode': True}), ('docx', {}), ('docx', {'plain_text_mode': True})])
async def test_resume_saved_draft_only_polishes_then_translates_next(tmp_path, provider, fmt, options):
    calls, _ = provider
    manager, output, kwargs = setup_job(tmp_path, fmt, **options)
    stopped = {'value': False}
    def saved(tid, index, source, translation):
        row = manager.db.get_chunks(tid)[-1]
        if row['chunk_data'].get('polish_status') == 'pending':
            stopped['value'] = True
    manager.db.passage_callback = saved
    kwargs['check_interruption_callback'] = lambda: stopped['value']
    assert not await translate_file(**kwargs)
    assert [call[0] for call in calls] == ['T']
    row = manager.db.get_chunks('job')[0]
    assert row['chunk_data']['polish_status'] == 'pending'
    calls.clear()
    manager.db.passage_callback = None
    stopped['value'] = False
    assert await translate_file(**kwargs)
    assert [call[0] for call in calls[:3]] == ['P', 'T', 'P']
    assert all(is_polished(row) for row in manager.db.get_chunks('job'))


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt,options', [('txt', {}), ('srt', {}), ('epub', {}), ('epub', {'plain_text_mode': True}), ('docx', {}), ('docx', {'plain_text_mode': True})])
async def test_failed_polish_keeps_draft_retry_never_retranslates(tmp_path, provider, fmt, options):
    calls, control = provider
    manager, output, kwargs = setup_job(tmp_path, fmt, **options)
    control['fail_polish'] = True
    assert not await translate_file(**kwargs)
    rows = manager.db.get_chunks('job')
    assert all(row['status'] == 'failed' and row['translated_text'] == row['chunk_data']['draft_text'] for row in rows)
    calls.clear()
    control['fail_polish'] = False
    assert await translate_file(**kwargs)
    assert [call[0] for call in calls] == ['P'] * len(rows)


@pytest.mark.asyncio
async def test_rate_limit_preserves_pending_draft(tmp_path, provider):
    calls, control = provider
    manager, _, kwargs = setup_job(tmp_path, 'docx')
    control['rate_limit'] = True
    with pytest.raises(RateLimitError):
        await translate_file(**kwargs)
    row = manager.db.get_chunks('job')[0]
    assert row['chunk_data']['polish_status'] == 'pending'
    calls.clear()
    control['rate_limit'] = False
    assert await translate_file(**kwargs)
    assert calls[0][0] == 'P'


@pytest.mark.asyncio
async def test_structured_corrupt_polish_keeps_draft(tmp_path, provider):
    _, control = provider
    manager, _, kwargs = setup_job(tmp_path, 'epub', structured_refinement_hide_placeholders=False)
    control['corrupt'] = True
    assert not await translate_file(**kwargs)
    assert all(row['chunk_data']['polish_status'] == 'failed' and row['translated_text'] == row['chunk_data']['draft_text'] for row in manager.db.get_chunks('job'))


@pytest.mark.asyncio
async def test_legacy_completed_rows_are_drafts_for_new_workflow(tmp_path, provider):
    calls, _ = provider
    manager, _, kwargs = setup_job(tmp_path, 'txt')
    assert await translate_file(**{**kwargs, 'prompt_options': {}})
    legacy = manager.db.get_chunks('job')
    assert all('draft_text' not in row['chunk_data'] for row in legacy)
    calls.clear()
    assert await translate_file(**kwargs)
    assert [call[0] for call in calls] == ['P'] * len(legacy)


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt', ['txt', 'srt', 'epub', 'docx'])
async def test_handler_uses_two_requests_and_stable_monotonic_progress(tmp_path, provider, fmt):
    from unittest.mock import Mock
    from src.api.handlers import perform_actual_translation
    from src.api.job_runner import JobRunner
    from src.api.translation_state import TranslationStateManager
    calls, _ = provider
    manager, output, kwargs = setup_job(tmp_path, fmt)
    state = TranslationStateManager(manager)
    config = dict(file_type=fmt, file_path=kwargs['input_filepath'], output_filename=output.name,
                  source_language='English', target_language='Korean', model='fake-echo',
                  llm_provider='poe', llm_api_endpoint='http://localhost:9999', refine_after=True,
                  prompt_options={}, max_tokens_per_chunk=60, auto_adjust_context=False,
                  srt_lines_per_block=2, connection_id='offline', is_resume=True)
    manager.update_job_config('job', config)
    state.create_translation('job', config)
    socket = Mock()
    runner = JobRunner(state, tmp_path, socket)
    runner.active = 'job'
    await perform_actual_translation('job', config, state, str(tmp_path), socket)
    assert state.get_translation_field('job', 'status') == 'completed', state.get_translation_field('job', 'error')
    rows = manager.db.get_chunks('job')
    assert [call[0] for call in calls] == ['T', 'P'] * len(rows)
    stats = state.get_translation_field('job', 'stats')
    assert stats['requests'] == 2 * len(rows)
    assert stats['completed_chunks'] == stats['total_chunks'] == len(rows)
    summaries = [call.args[1]['job'] for call in socket.emit.call_args_list if call.args[0] == 'job_update']
    percents = [summary['percent'] for summary in summaries]
    assert percents == sorted(percents) and percents[-1] == 100
    assert {summary['phase'] for summary in summaries} >= {'translating', 'polishing'}
    assert any(0 < p < 50 for p in percents)
    assert all(summary['units_total'] == len(rows) for summary in summaries if summary['units_total'])
    assert any(summary['current_unit'] == 0 for summary in summaries)
    assert max(summary['current_unit'] for summary in summaries if summary['current_unit'] is not None) == len(rows) - 1
    final_config = manager.get_job('job')['config']
    assert final_config['output_filepath'] == str(output)
    draft_path = final_config['refinement_source_path']
    if fmt in ('txt', 'srt'):
        from pathlib import Path
        assert 'DRAFT' in Path(draft_path).read_text() and 'POLISHED' not in Path(draft_path).read_text()
    assert not final_config['refinement_stale']


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt,options', [('txt', {}), ('srt', {}), ('epub', {}), ('epub', {'chapter_mode': True}), ('epub', {'plain_text_mode': True}), ('docx', {}), ('docx', {'plain_text_mode': True})])
async def test_corrective_replay_uses_drafts_updates_final_rows_and_revision(tmp_path, provider, fmt, options):
    from src.core.adapters.refine_file import refine_file
    calls, _ = provider
    manager, output, kwargs = setup_job(tmp_path, fmt, **options)
    assert await translate_file(**kwargs)
    rows = manager.db.get_chunks('job')
    config = manager.get_job('job')['config']
    config.update(_resync_original_file_path=kwargs['input_filepath'], source_language='English')
    manager.update_job_config('job', config)
    revision = manager.get_job('job')['passages_revision']
    manager.mark_refinement_stale('job')
    calls.clear()
    assert await refine_file(input_filepath=str(output), output_filepath=str(output), target_language='Korean',
                             model_name='fake-echo', llm_provider='poe', checkpoint_manager=manager,
                             translation_id='job', max_tokens_per_chunk=60, auto_adjust_context=False,
                             prompt_options={'_refine_after': True, **options}, srt_lines_per_block=2)
    assert [call[0] for call in calls] == ['P'] * len(rows)
    assert all('POLISHED' not in fake_llm._extract_source_block(call[2]) for call in calls)
    assert all(is_polished(row) for row in manager.db.get_chunks('job'))
    assert manager.get_job('job')['passages_revision'] > revision


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt,options', [('txt', {}), ('epub', {}), ('docx', {'plain_text_mode': True})])
async def test_source_analysis_once_and_same_snapshot_for_both_requests(tmp_path, provider, monkeypatch, fmt, options):
    import src.config as config
    import src.core.translator as translator
    import src.prompts.prompts as prompts
    from src.utils.novel_context import decode_context_snapshot
    analyses, translation_contexts, polish_contexts = [], [], []
    monkeypatch.setattr(config, 'NOVEL_CONTEXTS_DIR', tmp_path / 'notes')

    async def analyze(**kwargs):
        assert 'SOURCE' in kwargs['source_chunk'] and 'DRAFT' not in kwargs['source_chunk']
        analyses.append(kwargs['source_chunk'])
        return ('# GLOBAL LORE\n\n## GLOSSARY & TERMINOLOGY\n- traveller: fact ' + str(len(analyses)),
                'Source relationship ' + str(len(analyses)), [])
    monkeypatch.setattr('src.utils.novel_context.update_novel_context_chunk', analyze)
    original_translate = translator.generate_translation_prompt
    original_refine = translator.generate_refinement_prompt
    original_post = prompts.generate_post_processing_prompt

    def record_translate(*args, **kwargs):
        translation_contexts.append((kwargs['prompt_options']['novel_context'], kwargs['prompt_options'].get('dialogue_attribution')))
        return original_translate(*args, **kwargs)
    def record_refine(*args, **kwargs):
        polish_contexts.append((kwargs['prompt_options']['novel_context'], kwargs['prompt_options'].get('dialogue_attribution')))
        return original_refine(*args, **kwargs)
    def record_post(*args, **kwargs):
        polish_contexts.append((kwargs['prompt_options']['novel_context'], kwargs['prompt_options'].get('dialogue_attribution')))
        return original_post(*args, **kwargs)
    monkeypatch.setattr(translator, 'generate_translation_prompt', record_translate)
    monkeypatch.setattr(translator, 'generate_refinement_prompt', record_refine)
    monkeypatch.setattr(prompts, 'generate_post_processing_prompt', record_post)
    manager, _, kwargs = setup_job(tmp_path, fmt, auto_update_context=True, **options)
    stopped = {'value': False}
    manager.db.passage_callback = lambda *args: stopped.update(value=True)
    kwargs['check_interruption_callback'] = lambda: stopped['value']
    assert not await translate_file(**kwargs)
    assert len(analyses) == 1
    manager.db.passage_callback = None
    stopped['value'] = False
    assert await translate_file(**kwargs)
    rows = manager.db.get_chunks('job')
    assert len(analyses) == len(rows)
    assert translation_contexts == polish_contexts
    assert translation_contexts == [(decode_context_snapshot(row['chunk_data']['context_snapshot'], '')[0], row['chunk_data'].get('dialogue_attribution')) for row in rows]


@pytest.mark.asyncio
async def test_plain_segment_reassembly_keeps_empty_image_paragraph_and_long_paragraph(tmp_path, provider):
    from src.core.epub.epub_translation_adapter import EpubTranslationAdapter
    root = etree.fromstring(('<html><body><p>SOURCE0 ' + 'A long sentence. ' * 80
                            + '</p><p><img src="picture.png" /></p><p>SOURCE1 Last paragraph.</p></body></html>').encode())
    manager = _checkpoint_manager(tmp_path)
    manager.start_job('job', 'epub', {'refine_after': True})
    from src.core.llm_client import LLMClient
    ok, stats = await EpubTranslationAdapter().translate_content(
        root, {}, {}, 'English', 'Korean', 'fake', LLMClient(provider_type='poe'), 50,
        prompt_options={'_passage_workflow': 1, '_draft_artifacts': {}, 'plain_text_mode': True},
        checkpoint_manager=manager, translation_id='job', file_href='chapter.xhtml')
    assert ok and stats.total_chunks > 2
    paragraphs = root.findall('.//p')
    assert 'POLISHED0' in ''.join(paragraphs[0].itertext())
    assert paragraphs[1].find('img').get('src') == 'picture.png'
    assert 'POLISHED1' in ''.join(paragraphs[2].itertext())
    assert len(paragraphs) == 3


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt,options', [('txt', {}), ('epub', {}), ('docx', {}), ('docx', {'plain_text_mode': True})])
async def test_continuation_copies_final_draft_and_failed_stages_without_retranslation(tmp_path, provider, fmt, options):
    calls, _ = provider
    manager, _, kwargs = setup_job(tmp_path, fmt, **options)
    kwargs['max_tokens_per_chunk'] = 35
    assert await translate_file(**kwargs)
    base = manager.db.get_chunks('job')
    assert len(base) >= 3
    # A continuation may follow an interrupted job whose later drafts exist.
    pending = base[1]
    failed = base[2]
    for row, status in [(pending, 'pending'), (failed, 'failed')]:
        data = {**row['chunk_data'], 'polish_status': status}
        manager.db.save_chunk('job', row['chunk_index'], row['original_text'], data['draft_text'], data)
    manager.start_job('continued', fmt, {'refine_after': True, '_passage_workflow': 1}, kwargs['input_filepath'])
    calls.clear()
    kwargs['translation_id'] = 'continued'
    kwargs['continuation_base_id'] = 'job'
    assert await translate_file(**kwargs)
    assert calls and all(call[0] == 'P' for call in calls)
    continued = manager.db.get_chunks('continued')
    assert len(continued) == len(base) and all(is_polished(row) for row in continued)
    assert continued[0]['translated_text'] == base[0]['translated_text']
    assert continued[0]['chunk_data']['draft_text'] == base[0]['chunk_data']['draft_text']


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt', ['epub', 'docx'])
async def test_sqlite_final_rows_override_stale_structured_json_cache(tmp_path, provider, fmt):
    import json
    calls, _ = provider
    manager, output, kwargs = setup_job(tmp_path, fmt)
    assert await translate_file(**kwargs)
    states = list((manager.uploads_dir / 'job' / 'xhtml_states').glob('*.json'))
    assert states
    for path in states:
        state = json.loads(path.read_text())
        state['translated_chunks'] = ['STALE CACHE'] * len(state['translated_chunks'])
        path.write_text(json.dumps(state))
    calls.clear()
    assert await translate_file(**kwargs)
    assert calls == []
    with zipfile.ZipFile(output) as archive:
        assert not any(b'STALE CACHE' in archive.read(name) for name in archive.namelist())


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt', ['txt', 'srt', 'epub', 'docx'])
async def test_polish_only_rows_are_single_stage_and_use_passage_progress(tmp_path, provider, fmt):
    from unittest.mock import Mock
    from src.api.handlers import perform_actual_translation
    from src.api.job_runner import JobRunner
    from src.api.translation_state import TranslationStateManager
    calls, _ = provider
    manager, output, kwargs = setup_job(tmp_path, fmt)
    config = dict(file_type=fmt, file_path=kwargs['input_filepath'], output_filename=output.name,
                  source_language='Korean', target_language='Korean', model='fake-echo',
                  llm_provider='poe', llm_api_endpoint='http://localhost:9999', refine_only=True,
                  prompt_options={}, max_tokens_per_chunk=60, auto_adjust_context=False,
                  srt_lines_per_block=2, connection_id='offline', is_resume=True)
    manager.update_job_config('job', config)
    state = TranslationStateManager(manager)
    state.create_translation('job', config)
    socket = Mock()
    runner = JobRunner(state, tmp_path, socket)
    runner.active = 'job'
    await perform_actual_translation('job', config, state, str(tmp_path), socket)
    assert state.get_translation_field('job', 'status') == 'completed'
    rows = manager.db.get_chunks('job')
    assert len(calls) == len(rows) and all(call[0] == 'P' for call in calls)
    assert all(row['chunk_data']['draft_text'] == row['original_text'] and is_polished(row) for row in rows)
    stats = state.get_translation_field('job', 'stats')
    assert stats['total_chunks'] == stats['completed_chunks'] == len(rows)
    assert stats['requests'] == len(rows) and stats['percent'] == 100


def test_failed_translation_is_not_counted_as_a_finished_polish_decision():
    from src.core.common.passage_workflow import passage_progress
    progress = passage_progress([{'status': 'partial', 'chunk_data': {}, 'translated_text': 'source fallback'}], 2)
    assert progress['failed_chunks'] == 1
    assert progress['completed_chunks'] == 0 and progress['percent'] == 0


@pytest.mark.asyncio
async def test_notes_edit_during_polish_cannot_finalize_an_old_snapshot(tmp_path, provider, monkeypatch):
    from src.core.llm_client import LLMClient
    from src.core import translator
    from src.utils.novel_context import compress_dynamic_state
    manager = _checkpoint_manager(tmp_path)
    manager.start_job('job', 'txt', {'refine_after': True, '_passage_workflow': 1})
    old_snapshot = compress_dynamic_state('# GLOBAL LORE\n\n## GLOSSARY & TERMINOLOGY\n- Alice: old fact')
    new_snapshot = compress_dynamic_state('# GLOBAL LORE\n\n## GLOSSARY & TERMINOLOGY\n- Alice: corrected fact')
    async def request(**kwargs):
        row = manager.db.get_chunks('job')[0]
        manager.db.save_chunk('job', 0, row['original_text'], row['translated_text'], {'context_snapshot': new_snapshot})
        manager.mark_refinement_stale('job')
        return 'POLISHED stale result', None
    monkeypatch.setattr(translator, '_make_refinement_request', request)
    workflow = PassageWorkflow(manager, 'job', LLMClient(provider_type='poe'), 'fake', 'Korean', {})
    _, ok = await workflow.finish(0, 'SOURCE0', 'DRAFT0', {'context_snapshot': old_snapshot})
    row = manager.db.get_chunks('job')[0]
    assert not ok and row['chunk_data']['polish_status'] == 'pending'
    assert row['chunk_data']['context_snapshot'] == new_snapshot
    assert row['chunk_data']['draft_text'] == 'DRAFT0'


@pytest.mark.asyncio
async def test_resumed_txt_polish_restores_previous_polished_tail(tmp_path, provider):
    from src.core.llm_client import LLMClient
    calls, _ = provider
    manager = _checkpoint_manager(tmp_path)
    manager.start_job('job', 'txt', {'refine_after': True})
    manager.db.save_chunk('job', 0, 'SOURCE0', 'POLISHED0 UNIQUE_TAIL',
                          {'draft_text': 'DRAFT0', 'polish_status': 'done', 'passage_workflow_version': 1})
    manager.db.save_chunk('job', 1, 'SOURCE1', 'DRAFT1',
                          {'draft_text': 'DRAFT1', 'polish_status': 'pending', 'passage_workflow_version': 1})
    workflow = PassageWorkflow(manager, 'job', LLMClient(provider_type='poe'), 'fake', 'Korean', {})
    assert (await workflow.finish(1, 'SOURCE1', 'DRAFT1', {}))[1]
    assert calls[0][0] == 'P' and 'POLISHED0 UNIQUE_TAIL' in calls[0][2]
